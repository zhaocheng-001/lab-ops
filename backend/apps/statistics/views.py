"""基础统计视图：支持按实验室、设备类型、时间筛选。"""
import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from . import services


def _params(request):
    return request.query_params


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def overview(request):
    """总览：预约总数、待处理、使用中、已完成、逾期数、设备使用率。"""
    user = request.user
    params = _params(request)
    data = services.overall_statistics(user, params)
    data["status_distribution"] = services.status_distribution(user, params)
    return Response({"code": 0, "message": "success", "data": data})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def by_laboratory(request):
    return Response({"code": 0, "message": "success",
                     "data": services.by_laboratory(request.user, _params(request))})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def by_category(request):
    return Response({"code": 0, "message": "success",
                     "data": services.by_category(request.user, _params(request))})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def trend(request):
    group = request.query_params.get("group", "day")
    return Response({"code": 0, "message": "success",
                     "data": services.trend(request.user, _params(request), group)})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def device_ranking(request):
    limit = int(request.query_params.get("limit", 10))
    return Response({"code": 0, "message": "success",
                     "data": services.device_ranking(request.user, _params(request), limit)})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def applicant_ranking(request):
    limit = int(request.query_params.get("limit", 10))
    return Response({"code": 0, "message": "success",
                     "data": services.applicant_ranking(request.user, _params(request), limit)})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def dashboard(request):
    """统计看板：一次性返回多个统计维度，减少前端请求次数。"""
    user = request.user
    params = _params(request)
    return Response({
        "code": 0,
        "message": "success",
        "data": {
            "overview": services.overall_statistics(user, params),
            "status_distribution": services.status_distribution(user, params),
            "by_laboratory": services.by_laboratory(user, params),
            "by_category": services.by_category(user, params),
            "trend": services.trend(user, params, params.get("group", "day")),
            "device_ranking": services.device_ranking(user, params, 10),
            "applicant_ranking": services.applicant_ranking(user, params, 10),
        },
    })


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def export_report(request):
    """导出统计报表（Excel）。"""
    user = request.user
    params = _params(request)
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "统计总览"
    header_fill = PatternFill("solid", fgColor="4472C4")
    sheet.append(["统计指标", "数值"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    overview_data = services.overall_statistics(user, params)
    label_map = {
        "total": "预约总数", "pending": "待处理", "in_use": "使用中", "completed": "已完成",
        "overdue": "逾期数", "rejected": "已驳回", "cancelled": "已取消",
        "device_usage_rate": "设备使用率(%)", "completion_rate": "完成率(%)",
        "used_hours": "累计使用时长(小时)", "average_score": "平均评分",
    }
    for key, label in label_map.items():
        sheet.append([label, overview_data.get(key, 0)])
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 16

    sheet2 = workbook.create_sheet("按实验室统计")
    sheet2.append(["实验室", "位置", "预约总数", "待审批", "使用中", "已完成", "逾期", "使用时长(小时)"])
    for cell in sheet2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in services.by_laboratory(user, params):
        sheet2.append([row["laboratory"], row["location"], row["total"], row["pending"],
                       row["in_use"], row["completed"], row["overdue"], row["hours"]])
    for column, width in zip("ABCDEFGH", [22, 20, 12, 12, 12, 12, 10, 16]):
        sheet2.column_dimensions[column].width = width

    sheet3 = workbook.create_sheet("按设备类型统计")
    sheet3.append(["设备类型", "预约总数", "已完成", "逾期", "使用时长(小时)"])
    for cell in sheet3[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in services.by_category(user, params):
        sheet3.append([row["category"], row["total"], row["completed"], row["overdue"], row["hours"]])
    for column, width in zip("ABCDE", [22, 12, 12, 10, 16]):
        sheet3.column_dimensions[column].width = width

    sheet4 = workbook.create_sheet("预约趋势")
    sheet4.append(["时间", "预约总数", "已完成", "已驳回", "已取消"])
    for cell in sheet4[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in services.trend(user, params, params.get("group", "day")):
        sheet4.append([row["period"], row["total"], row["completed"],
                       row["rejected"], row["cancelled"]])

    sheet5 = workbook.create_sheet("设备使用排行")
    sheet5.append(["设备", "编号", "类型", "实验室", "预约次数", "使用时长(小时)", "使用率(%)"])
    for cell in sheet5[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in services.device_ranking(user, params, 50):
        sheet5.append([row["device"], row["code"], row["category"], row["laboratory"],
                       row["total"], row["hours"], row["usage_rate"]])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        "attachment; filename=statistics.xlsx; filename*=UTF-8''统计报表.xlsx"
    )
    return response
